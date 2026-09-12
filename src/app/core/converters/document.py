import csv
from pathlib import Path

from docx import Document
from openpyxl import Workbook, load_workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

try:
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    CJK_FONT = "STSong-Light"
except Exception:
    CJK_FONT = "Helvetica"


def convert_document(source, output):
    source = Path(source)
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)

    source_ext = source.suffix.lower().lstrip(".")
    target_ext = output.suffix.lower().lstrip(".")

    if source_ext == "docx" and target_ext in {"txt", "md", "html"}:
        _docx_to_text(source, output, target_ext)
    elif source_ext == "xlsx" and target_ext == "csv":
        _xlsx_to_csv(source, output)
    elif source_ext == "csv" and target_ext == "xlsx":
        _csv_to_xlsx(source, output)
    elif source_ext in {"txt", "md"} and target_ext == "pdf":
        _text_or_markdown_to_pdf(source, output, source_ext)
    elif source_ext == "csv" and target_ext == "pdf":
        _csv_to_pdf(source, output)
    elif source_ext in {"html", "htm"} and target_ext == "pdf":
        _html_to_pdf(source, output)
    else:
        raise ValueError(f"暂不支持 .{source_ext} 转换到 .{target_ext}")


def _docx_to_text(source, output, target_ext):
    document = Document(source)
    paragraphs = []
    for paragraph in document.paragraphs:
        text = paragraph.text
        if paragraph.style.name.startswith("Heading"):
            level = paragraph.style.name.replace("Heading", "").strip() or "1"
            try:
                level = int(level)
            except ValueError:
                level = 1
            if target_ext == "md":
                paragraphs.append(f"{'#' * max(1, min(6, level))} {text}")
            elif target_ext == "html":
                paragraphs.append(f"<h{max(1, min(6, level))}>{text}</h{max(1, min(6, level))}>")
            else:
                paragraphs.append(text)
        else:
            if target_ext == "html":
                paragraphs.append(f"<p>{text}</p>")
            else:
                paragraphs.append(text)

    if target_ext == "html":
        content = "\n".join(paragraphs)
    else:
        content = "\n\n".join(paragraphs)
    output.write_text(content, encoding="utf-8")


def _xlsx_to_csv(source, output):
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        with output.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle)
            for row in sheet.iter_rows(values_only=True):
                writer.writerow(["" if value is None else value for value in row])
    finally:
        workbook.close()


def _csv_to_xlsx(source, output):
    workbook = Workbook()
    sheet = workbook.active
    with source.open("r", encoding="utf-8-sig", newline="") as handle:
        for row in csv.reader(handle):
            sheet.append([_coerce_scalar(value) for value in row])
    workbook.save(output)


def _coerce_scalar(value):
    value = value.strip()
    try:
        return int(value)
    except ValueError:
        pass
    try:
        return float(value)
    except ValueError:
        return value


def _text_or_markdown_to_pdf(source, output, source_ext):
    text = source.read_text(encoding="utf-8")
    if source_ext == "md":
        import markdown

        html = markdown.markdown(text)
        _html_to_pdf_from_text(html, output)
        return
    _write_paragraphs(text.splitlines(), output)


def _csv_to_pdf(source, output):
    with source.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = [row for row in csv.reader(handle) if row]
    _write_table(rows, output)


def _html_to_pdf(source, output):
    html = source.read_text(encoding="utf-8")
    _html_to_pdf_from_text(html, output)


def _html_to_pdf_from_text(html, output):
    from xhtml2pdf import pisa

    with output.open("wb") as handle:
        pisa.CreatePDF(html, dest=handle)


def _write_paragraphs(lines, output):
    styles = getSampleStyleSheet()
    style = ParagraphStyle(
        "ChineseBody",
        parent=styles["BodyText"],
        fontName=CJK_FONT,
        fontSize=11,
        leading=16,
    )
    document = SimpleDocTemplate(str(output), pagesize=A4)
    story = []
    for line in lines:
        line = line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        story.append(Paragraph(line or "&nbsp;", style))
        story.append(Spacer(1, 2 * mm))
    document.build(story)


def _write_table(rows, output):
    styles = getSampleStyleSheet()
    style = ParagraphStyle(
        "TableText",
        parent=styles["BodyText"],
        fontName=CJK_FONT,
        fontSize=10,
        leading=13,
    )
    data = []
    for row in rows:
        data.append([Paragraph(_escape_xml(str(cell)), style) for cell in row])

    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    document = SimpleDocTemplate(str(output), pagesize=A4)
    document.build([table])


def _escape_xml(value):
    return value.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
