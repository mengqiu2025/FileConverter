import csv
import sys
import tempfile
import unittest
from pathlib import Path

SRC = Path(__file__).resolve().parents[1] / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from docx import Document
from openpyxl import load_workbook
import pymupdf as fitz

from app.core.converters.document import convert_document


class DocumentConverterTests(unittest.TestCase):
    def test_docx_to_txt_extracts_text(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            source = tmp / "input.docx"
            output = tmp / "converted" / "input.txt"
            document = Document()
            document.add_paragraph("HELLO DOCX")
            document.save(source)

            convert_document(source, output)

            self.assertTrue(output.exists())
            self.assertIn("HELLO DOCX", output.read_text(encoding="utf-8"))

    def test_xlsx_to_csv_extracts_rows(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            source = tmp / "input.xlsx"
            output = tmp / "converted" / "input.csv"
            from openpyxl import Workbook

            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["name", "score"])
            sheet.append(["Ada", 95])
            workbook.save(source)

            convert_document(source, output)

            self.assertTrue(output.exists())
            with output.open("r", encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.reader(handle))
            self.assertEqual(rows[0], ["name", "score"])
            self.assertEqual(rows[1], ["Ada", "95"])

    def test_csv_to_xlsx_creates_workbook(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            source = tmp / "input.csv"
            output = tmp / "converted" / "input.xlsx"
            source.write_text("name,score\nAda,95\n", encoding="utf-8")

            convert_document(source, output)

            self.assertTrue(output.exists())
            workbook = load_workbook(output)
            sheet = workbook.active
            self.assertEqual(sheet["A1"].value, "name")
            self.assertEqual(sheet["B2"].value, 95)

    def test_txt_to_pdf_creates_pdf_with_text(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            source = tmp / "input.txt"
            output = tmp / "converted" / "input.pdf"
            source.write_text("HELLO TXT", encoding="utf-8")

            convert_document(source, output)

            self.assertTrue(output.exists())
            document = fitz.open(output)
            try:
                text = document[0].get_text()
            finally:
                document.close()
            self.assertIn("HELLO TXT", text)

    def test_csv_to_pdf_creates_table_pdf(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            source = tmp / "input.csv"
            output = tmp / "converted" / "input.pdf"
            source.write_text("name,score\nAda,95\n", encoding="utf-8")

            convert_document(source, output)

            self.assertTrue(output.exists())
            document = fitz.open(output)
            try:
                text = document[0].get_text()
            finally:
                document.close()
            self.assertIn("name", text)
            self.assertIn("Ada", text)

    def test_html_to_pdf_creates_pdf(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            source = tmp / "input.html"
            output = tmp / "converted" / "input.pdf"
            source.write_text("<h1>HELLO HTML</h1><p>World</p>", encoding="utf-8")

            convert_document(source, output)

            self.assertTrue(output.exists())
            document = fitz.open(output)
            try:
                text = document[0].get_text()
            finally:
                document.close()
            self.assertIn("HELLO HTML", text)
            self.assertIn("World", text)


if __name__ == "__main__":
    unittest.main()
